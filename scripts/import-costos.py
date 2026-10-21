#!/usr/bin/env python3
"""
Importador de los catálogos de costos de Cobalto Barroco.

Lee los Excel del flujo de presupuestos y produce un JSON con los catálogos
GLOBALES del módulo de Precios Unitarios, listos para la base de datos de la
plataforma (claves `costMaterials`, `recipes`, `laborRates`):

  - ` MATERIALES`  → costMaterials  (material / herramienta / equipo + precio + url)
  - `RECETAS`      → recipes         (soluciones compuestas con su BOM)
  - `SALARIOS`     → laborRates      (tarifa de mano de obra por rol)

Uso:
    python import-costos.py --excel-dir "C:/.../Desktop" --out costos-seed.json
    python import-costos.py --merge "C:/.../.localdata/cobalto.json"

`--merge` inyecta los catálogos en una base de datos existente (reemplaza esas
tres colecciones y conserva el resto). Sin `--merge`, solo escribe el seed.

Tipos: ver `src/lib/db.ts` (CostMaterial, Recipe, LaborRate).
"""
import argparse
import json
import os
import re
import sys
import time

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instala con: pip install openpyxl")

NOW = int(time.time() * 1000)


def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip()).upper()


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def find_excel(excel_dir, needle):
    for n in os.listdir(excel_dir):
        if needle in n and n.lower().endswith(".xlsx"):
            return os.path.join(excel_dir, n)
    raise SystemExit(f"No encontré el Excel que contiene '{needle}' en {excel_dir}")


def cell(row, idx1):
    """Valor de la columna idx1 (1-indexed) de una tupla de fila, o None."""
    i = idx1 - 1
    return row[i] if 0 <= i < len(row) else None


# ── MATERIALES ──────────────────────────────────────────────────────────────
# Tres bloques en paralelo; por bloque: (col_url, col_nombre, col_precio).
MAT_BLOCKS = [
    ("material", 1, 2, 3),
    ("herramienta", 5, 6, 7),
    ("equipo", 9, 10, 11),
]
MAT_HEADERS = {"MATERIAL", "HERRAMIENTA", "EQUIPO", "DESCRIPCIÓN", "P.U.", "PU"}


def import_materiales(ws):
    out = []
    seq = 0
    cat = {b[0]: "" for b in MAT_BLOCKS}  # categoría vigente por bloque
    for row in ws.iter_rows(min_row=3, values_only=True):
        for kind, c_url, c_name, c_price in MAT_BLOCKS:
            name = cell(row, c_name)
            price = cell(row, c_price)
            if name is None or str(name).strip() == "":
                continue
            name = str(name).strip()
            if norm(name) in MAT_HEADERS:
                continue
            if is_num(price):
                seq += 1
                url = cell(row, c_url)
                url = str(url).strip() if isinstance(url, str) and url.strip().lower().startswith("http") else ""
                out.append(
                    {
                        "id": f"MAT-{seq:04d}",
                        "nombre": name,
                        "kind": kind,
                        "categoria": cat[kind],
                        "precio": round(float(price), 4),
                        "unidad": "",
                        "url": url,
                        "supplierId": None,
                        "nota": "",
                        "actualizadoEn": NOW,
                        "createdAt": NOW,
                    }
                )
            else:
                # Nombre sin precio numérico = subtítulo de categoría del bloque.
                cat[kind] = name
    return out


# ── RECETAS ─────────────────────────────────────────────────────────────────
def import_recetas(ws):
    recipes = []
    cur = None
    for row in ws.iter_rows(min_row=3, values_only=True):
        sol = cell(row, 2)
        mat = cell(row, 3)
        if isinstance(sol, str) and sol.strip():
            cur = {
                "id": f"REC-{len(recipes) + 1:04d}",
                "nombre": sol.strip(),
                "rendimiento": str(cell(row, 10) or "").strip(),
                "unidad": "",
                "componentes": [],
                "costo": 0.0,
                "costoFinal": round(float(cell(row, 9)), 4) if is_num(cell(row, 9)) else None,
                "nota": "",
                "createdAt": NOW,
            }
            recipes.append(cur)
        elif cur is not None and isinstance(mat, str) and mat.strip():
            precio = cell(row, 5)
            cant = cell(row, 6)
            cur["componentes"].append(
                {
                    "materialId": None,
                    "recipeId": None,
                    "nombre": mat.strip(),
                    "unidad": str(cell(row, 4) or "").strip(),
                    "precio": round(float(precio), 4) if is_num(precio) else 0.0,
                    "cantidad": round(float(cant), 6) if is_num(cant) else 0.0,
                }
            )
    # costo = Σ precio×cantidad; derivar unidad del rendimiento ("1 L" → "L")
    for r in recipes:
        r["costo"] = round(sum(c["precio"] * c["cantidad"] for c in r["componentes"]), 4)
        m = re.search(r"([A-Za-zÁÉÍÓÚÑ]+)\s*$", r["rendimiento"])
        r["unidad"] = m.group(1) if m else ""
    return recipes


def link_recipe_refs(recipes, materials):
    """Resuelve cada componente a un material o receta del catálogo por nombre."""
    mat_by_name = {norm(m["nombre"]): m["id"] for m in materials}
    rec_by_name = {norm(r["nombre"]): r["id"] for r in recipes}
    linked = 0
    for r in recipes:
        for c in r["componentes"]:
            key = norm(c["nombre"])
            if key in rec_by_name and rec_by_name[key] != r["id"]:
                c["recipeId"] = rec_by_name[key]
                linked += 1
            elif key in mat_by_name:
                c["materialId"] = mat_by_name[key]
                linked += 1
    return linked


# ── SALARIOS ─────────────────────────────────────────────────────────────────
def import_salarios(ws):
    out = []
    seq = 0
    for row in ws.iter_rows(min_row=4, values_only=True):
        rol = cell(row, 1)
        jornada = cell(row, 2)
        if not (isinstance(rol, str) and rol.strip()) or not is_num(jornada):
            continue
        seq += 1
        out.append(
            {
                "id": f"SAL-{seq:04d}",
                "rol": rol.strip(),
                "jornada": float(jornada),
                "hora": float(cell(row, 4)) if is_num(cell(row, 4)) else 0.0,
                "semana": float(cell(row, 5)) if is_num(cell(row, 5)) else 0.0,
                "mes": float(cell(row, 6)) if is_num(cell(row, 6)) else 0.0,
                "createdAt": NOW,
            }
        )
    return out


def sheet(wb, name):
    # La hoja de materiales viene con un espacio inicial (" MATERIALES").
    for s in wb.sheetnames:
        if s.strip().upper() == name.strip().upper():
            return wb[s]
    raise SystemExit(f"No hay hoja '{name}'. Hojas: {wb.sheetnames}")


def main():
    ap = argparse.ArgumentParser(description="Importa catálogos de costos a JSON.")
    ap.add_argument("--excel-dir", default=os.path.join(os.path.expanduser("~"), "Desktop"))
    ap.add_argument("--presupuesto", help="Ruta directa al Excel de presupuesto (opcional).")
    ap.add_argument("--out", default="costos-seed.json")
    ap.add_argument("--merge", help="Ruta al cobalto.json donde inyectar los catálogos.")
    args = ap.parse_args()

    f = args.presupuesto or find_excel(args.excel_dir, "PRESUPUESTO")
    print(f"Leyendo {f}")
    wb = openpyxl.load_workbook(f, data_only=True, read_only=True)

    materials = import_materiales(sheet(wb, "MATERIALES"))
    recipes = import_recetas(sheet(wb, "RECETAS"))
    labor = import_salarios(sheet(wb, "SALARIOS"))
    wb.close()

    linked = link_recipe_refs(recipes, materials)
    seed = {"costMaterials": materials, "recipes": recipes, "laborRates": labor}

    print(
        f"  materiales: {len(materials)}  "
        f"(material={sum(1 for m in materials if m['kind']=='material')}, "
        f"herramienta={sum(1 for m in materials if m['kind']=='herramienta')}, "
        f"equipo={sum(1 for m in materials if m['kind']=='equipo')})"
    )
    print(f"  recetas: {len(recipes)}  (componentes enlazados a catálogo: {linked})")
    print(f"  salarios: {len(labor)}")

    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(seed, fh, ensure_ascii=False, indent=2)
    print(f"Seed escrito en {args.out}")

    if args.merge:
        with open(args.merge, encoding="utf-8") as fh:
            db = json.load(fh)
        db["costMaterials"] = materials
        db["recipes"] = recipes
        db["laborRates"] = labor
        db.setdefault("concepts", [])
        db.setdefault("budgets", [])
        tmp = args.merge + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(db, fh, ensure_ascii=False, indent=2)
        os.replace(tmp, args.merge)
        print(f"Catálogos inyectados en {args.merge}")


if __name__ == "__main__":
    main()
