#!/usr/bin/env python3
"""Lee el Excel CURRÍCULUM (hoja única) y emite {cvEntries:[...]} para
importar en la plataforma vía el botón "Importar datos (JSON)" de Configuración.

Uso:  python3 -I scripts/import-cv.py <archivo.xlsx>  > cv.json

Estructura de la hoja (según el machote):
  A: CLAVE (CB001…, opcional)
  B: FECHA (mes/rango)  — o un AÑO (p. ej. "2026") en las filas de encabezado
  C: CLASIFICACIÓN (HISTÓRICO/MODERNO)
  D: TIPO DE BIEN
  E: LUGAR
  F: INMUEBLE/SITIO
  G: TÍTULO
  H: DESCRIPCIÓN
  I: COLABORACIÓN
"""
import sys
import json
import hashlib
import openpyxl


def hid(*parts):
    h = hashlib.md5("|".join(str(p) for p in parts).encode("utf-8")).hexdigest()
    return "CV-" + h[:12].upper()


def cell(ws, r, c):
    v = ws.cell(r, c).value
    if v is None:
        return ""
    return str(v).strip()


def main():
    if len(sys.argv) < 2:
        print("Falta el archivo .xlsx", file=sys.stderr)
        sys.exit(1)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ws = wb.worksheets[0]
    entries = []
    anio = None
    for r in range(1, ws.max_row + 1):
        A = cell(ws, r, 1)   # clave
        B = cell(ws, r, 2)   # fecha / año
        C = cell(ws, r, 3)
        D = cell(ws, r, 4)
        E = cell(ws, r, 5)
        F = cell(ws, r, 6)
        G = cell(ws, r, 7)
        H = cell(ws, r, 8)
        I = cell(ws, r, 9)
        # Fila de año: B es un número de 4 dígitos y el resto está vacío.
        if B.isdigit() and len(B) == 4 and not any([A, C, D, E, F, G, H, I]):
            anio = int(B)
            continue
        # Fila de datos: requiere título y estar bajo un año (ignora encabezado).
        if not G or anio is None:
            continue
        entries.append({
            "id": hid(A, G, anio, E),
            "clave": A or None,
            "anio": anio or 0,
            "fecha": B,
            "clasificacion": C or None,
            "tipoBien": D or None,
            "lugar": E or None,
            "sitio": F or None,
            "titulo": G,
            "descripcion": H or None,
            "colaboracion": I or None,
            "projectId": None,
            "createdAt": 0,
        })
    json.dump({"cvEntries": entries}, sys.stdout, ensure_ascii=False, indent=2)
    print(f"\n# {len(entries)} entradas de CV", file=sys.stderr)


if __name__ == "__main__":
    main()
